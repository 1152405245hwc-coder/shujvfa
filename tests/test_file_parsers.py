import io
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader, PdfWriter

from legal_funds_agent.parsers.file_parsers import extract_document_text, extract_transactions_csv, extract_transactions_csv_detailed


class FileParserTest(unittest.TestCase):
    def test_docx_text_extraction(self):
        document = Document()
        document.add_paragraph("起诉书：何某收取被害人朱某人民币50000元。")
        data = io.BytesIO()
        document.save(data)
        self.assertIn("起诉书", extract_document_text(data.getvalue(), filename="indictment.docx"))

    def test_pdf_text_extraction_reports_scanned_pdf(self):
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        data = io.BytesIO()
        writer.write(data)
        with self.assertRaisesRegex(ValueError, "OCR"):
            extract_document_text(data.getvalue(), filename="scan.pdf")

    def test_xlsx_bank_sheet_to_canonical_csv(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "01_账户"
        sheet.append(["户名", "何某", "账户ID", "A002"])
        sheet.append([None])
        sheet.append(["交易时间", "银行流水号", "收/支", "对方户名/账户", "金额", "摘要/备注"])
        sheet.append(["2025-03-12 10:21:14", "S001", "收入", "朱某", 50000, "项目款"])
        data = io.BytesIO()
        workbook.save(data)
        csv_text = extract_transactions_csv(data.getvalue(), filename="bank.xlsx")
        self.assertIn("S001,2025-03-12,10:21:14,朱某,,何某,A002,50000.00,项目款", csv_text)

    def test_case_package_formats_extract(self):
        root = Path(__file__).resolve().parents[1] / "sample_data" / "case_packages" / "GOLD_CASE_001"
        indictment = extract_document_text((root / "visible/documents/01_起诉书.docx").read_bytes(), filename="01_起诉书.docx")
        csv_text = extract_transactions_csv((root / "visible/bank/02_银行流水账单.xlsx").read_bytes(), filename="02_银行流水账单.xlsx")
        self.assertIn("起 诉 书", indictment)
        self.assertGreaterEqual(csv_text.count("\n"), 10)

    def test_gbk_csv_decoding_fallback(self):
        csv_text = "交易时间,银行流水号,收/支,对方户名/账户,金额,摘要/备注\n2025-03-12 10:00:00,TX1,收入,张某,50000,测试"
        gbk_bytes = csv_text.encode("gbk")
        text, stats = extract_transactions_csv_detailed(gbk_bytes, filename="bank.csv")
        self.assertIn("收入", text)
        self.assertEqual(stats["invalid_direction"], 0)
        self.assertEqual(stats["invalid_datetime"], 0)
        self.assertEqual(stats["invalid_amount_or_counterparty"], 0)

    def test_unrecognized_encoding_raises_value_error(self):
        data = b"\xff\xfe\x00\x01"
        with self.assertRaisesRegex(ValueError, "无法识别文件编码"):
            extract_transactions_csv_detailed(data, filename="bank.csv")

    def test_xlsx_skip_stats_count_invalid_rows(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "01_账户"
        sheet.append(["户名", "何某", "账户ID", "A002"])
        sheet.append([None])
        sheet.append(["交易时间", "银行流水号", "收/支", "对方户名/账户", "金额", "摘要/备注"])
        sheet.append(["2025-03-12 10:21:14", "S001", "收入", "朱某", 50000, "项目款"])
        sheet.append(["2025-03-12 10:21:15", "S002", "其他", "王某", 10000, ""])
        sheet.append(["不是日期", "S003", "支出", "陈某", 20000, ""])
        sheet.append(["2025-03-12 10:21:16", "S004", "支出", "", 30000, ""])
        sheet.append(["2025-03-12 10:21:17", "S005", "支出", "李某", None, ""])
        data = io.BytesIO()
        workbook.save(data)
        csv_text, stats = extract_transactions_csv_detailed(data.getvalue(), filename="bank.xlsx")
        self.assertIn("S001", csv_text)
        self.assertNotIn("S002", csv_text)
        self.assertNotIn("S003", csv_text)
        self.assertNotIn("S004", csv_text)
        self.assertNotIn("S005", csv_text)
        self.assertEqual(stats["invalid_direction"], 1)
        self.assertEqual(stats["invalid_datetime"], 1)
        self.assertEqual(stats["invalid_amount_or_counterparty"], 2)

    def test_screenshot_defaults_are_empty_not_fabricated(self):
        fake_parsed = {
            "transaction_id": "TX-IMG-123",
            "amount": "50000.00",
            "payee": "李某",
            "time": "",
        }
        with patch("legal_funds_agent.parsers.ocr_service.extract_text_from_image", return_value="mock"):
            with patch("legal_funds_agent.parsers.ocr_service.parse_screenshot_transaction", return_value=fake_parsed):
                csv_text = extract_transactions_csv(b"dummy", filename="screenshot.png")
        self.assertIn("TX-IMG-123,,,,,李某,,50000.00,转账截图识别", csv_text)
        self.assertNotIn("2026-03-15", csv_text)
        self.assertNotIn("12:00:00", csv_text)
        self.assertNotIn("被害人", csv_text)


if __name__ == "__main__":
    unittest.main()
