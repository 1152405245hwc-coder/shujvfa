# -*- coding: utf-8 -*-
import docx, os
D = r'C:\Users\youxi\Documents\shujvfa\sample_data\case_packages\GOLD_CASE_002\visible\documents'
checks = {
    '01_起诉书.docx': ['北京市某区人民检察院', '京某检刑诉〔2026〕87号', '此致', '检察官：王某某', '附：'],
    '03_刘某陈述.docx': ['被害人陈述笔录', '询问时间', '被询问人核对意见', '360万' if False else '240万元'],
    '04_周某陈述.docx': ['被害人陈述笔录', '被询问人核对意见'],
    '05_郑某陈述.docx': ['被害人陈述笔录', '被询问人核对意见'],
    '06_证人证言.docx': ['证 人 证 言 汇 编', '一、证人孙某证言', '七、证人马某甲证言'],
    '07_被告人供述与辩解.docx': ['第一次讯问', '第三次讯问', '讯问时间'],
}
ok = True
for name, keys in checks.items():
    d = docx.Document(os.path.join(D, name))
    text = '\n'.join(p.text for p in d.paragraphs) + '\n' + '\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
    miss = [k for k in keys if k not in text]
    sec = d.sections[0]
    info = 'A4' if abs(sec.page_width.cm - 21.0) < 0.1 else 'NOT-A4'
    print(name, info, 'OK' if not miss else ('MISSING: ' + str(miss)))
    if miss: ok = False
print('ALL', 'OK' if ok else 'FAIL')
