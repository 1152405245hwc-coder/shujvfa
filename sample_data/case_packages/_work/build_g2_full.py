# -*- coding: utf-8 -*-
"""
Rebuild GOLD_CASE_002 documents to the FULL standard format of GOLD_CASE_001:
- 01 起诉书: complete indictment (机关名称/文号/被告人段/案件来源/审查查明/证据/本院认为/此致/落款/附项)
- 03/04/05 victim statements: standard 询问笔录 (机关/标题/被害人/信息表/问与答/核对意见/签名)
- 06 witness statements: 证人证言汇编 with per-witness info table + Q&A
- 07 defendant statements: 讯问笔录 with per-session info table + Q&A
All facts (names, dates, amounts, summaries, qualifiers) preserved from originals.
Procedural details (询问时间/地点/民警/年龄性别等) are fictional-completed in GOLD_1 style.
"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = "C:/Users/youxi/Documents/shujvfa/sample_data/case_packages"
BUILD = os.path.join(ROOT, "_work", "build2")
os.makedirs(BUILD, exist_ok=True)

F_TITLE_EA = "黑体"
F_TITLE_ASCII = "SimHei"
F_BODY_EA = "仿宋"
F_BODY_ASCII = "Times New Roman"
F_NOTE_EA = "楷体"


def set_run_font(run, ea, ascii_f, size, bold=False):
    run.font.name = ascii_f
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), ascii_f)
    rFonts.set(qn('w:hAnsi'), ascii_f)
    rFonts.set(qn('w:eastAsia'), ea)


def set_first_line_chars(p, chars=200, twips=640):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), str(chars))
    ind.set(qn('w:firstLine'), str(twips))


def clear_indent(p):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), '0')
    ind.set(qn('w:firstLine'), '0')
    ind.set(qn('w:leftChars'), '0')
    ind.set(qn('w:left'), '0')


def para_props(p, align=None, line_exact=None, line_mult=None, before=0, after=0):
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if line_exact:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line_exact)
    elif line_mult:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line_mult
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def add_para(doc, text, kind, after=0):
    p = doc.add_paragraph()
    if kind == "title":
        para_props(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_mult=1.0, after=after)
        r = p.add_run(text)
        set_run_font(r, F_TITLE_EA, F_TITLE_ASCII, 22)
    elif kind == "subtitle":
        para_props(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_mult=1.0, after=after)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "h1":
        para_props(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_exact=28)
        set_first_line_chars(p)
        r = p.add_run(text)
        set_run_font(r, F_TITLE_EA, F_TITLE_ASCII, 16)
    elif kind == "qa":
        para_props(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_exact=28)
        set_first_line_chars(p)
        m = text[:2]
        r1 = p.add_run(m)
        set_run_font(r1, F_TITLE_EA, F_TITLE_ASCII, 16)
        r2 = p.add_run(text[2:])
        set_run_font(r2, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "body":
        para_props(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_exact=28)
        set_first_line_chars(p)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "thiszhi":
        para_props(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_exact=28)
        set_first_line_chars(p)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "court":
        para_props(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_exact=28)
        clear_indent(p)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "sign":
        para_props(p, align=WD_ALIGN_PARAGRAPH.RIGHT, line_exact=28)
        clear_indent(p)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "attach":
        para_props(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_exact=28)
        clear_indent(p)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "note":
        para_props(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_exact=24, before=6)
        set_first_line_chars(p, 200, 560)
        r = p.add_run(text)
        set_run_font(r, F_NOTE_EA, F_BODY_ASCII, 14)
    else:
        para_props(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_exact=28)
        set_first_line_chars(p)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    return p


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def vcenter_cell(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    v = OxmlElement('w:vAlign')
    v.set(qn('w:val'), 'center')
    tcPr.append(v)


def add_table(doc, rows_data, widths=None):
    n_rows = len(rows_data)
    n_cols = len(rows_data[0])
    if widths is None:
        if n_cols == 2:
            widths = [2.6, 13.0]
        elif n_cols == 4:
            widths = [2.6, 5.2, 2.6, 5.2]
        else:
            widths = [15.6 / n_cols] * n_cols
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for ri, row in enumerate(rows_data):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.width = Cm(widths[ci])
            vcenter_cell(cell)
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.space_before = Pt(1)
            pf.space_after = Pt(1)
            clear_indent(p)
            is_label = (ci % 2 == 0)
            is_short = len(val) <= 8
            if is_label:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shade_cell(cell, "F2F2F2")
                r = p.add_run(val)
                set_run_font(r, F_TITLE_EA, F_TITLE_ASCII, 14)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_short else WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(val)
                set_run_font(r, F_BODY_EA, F_BODY_ASCII, 14)
    return table


def new_doc():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = F_BODY_ASCII
    st.font.size = Pt(16)
    rPr = st.element.get_or_add_rPr()
    rF = rPr.find(qn('w:rFonts'))
    if rF is None:
        rF = OxmlElement('w:rFonts')
        rPr.append(rF)
    rF.set(qn('w:ascii'), F_BODY_ASCII)
    rF.set(qn('w:hAnsi'), F_BODY_ASCII)
    rF.set(qn('w:eastAsia'), F_BODY_EA)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3.7)
    sec.bottom_margin = Cm(3.5)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)
    return doc


def build(blocks, out):
    doc = new_doc()
    for idx, (kind, payload) in enumerate(blocks):
        if kind == "skip":
            continue
        if kind == "table":
            add_table(doc, payload)
            sp = doc.add_paragraph()
            para_props(sp, line_exact=14)
            continue
        after = 0
        if kind == "title":
            nxt = blocks[idx + 1][0] if idx + 1 < len(blocks) else None
            after = 6 if nxt == "title" else 18
        elif kind == "subtitle":
            after = 12
        add_para(doc, payload, kind, after=after)
    doc.save(out)


def qa(q, a):
    return [("qa", "问：" + q), ("qa", "答：" + a)]


# ============ 01 起诉书 ============
indictment = [
    ("title", "北京市某区人民检察院"),
    ("title", "起 诉 书"),
    ("subtitle", "京某检刑诉〔2026〕87号"),
    ("body", "被告人陈某，男，1987年出生，公民身份号码已隐去，汉族，大学文化，无固定职业，户籍所在地北京市某区。因涉嫌诈骗罪，于2025年8月26日被北京市公安局某分局刑事拘留；同年9月2日经本院批准，被依法逮捕。现羁押于北京市某区看守所。"),
    ("body", "本案由北京市公安局某分局侦查终结，以被告人陈某涉嫌诈骗罪，于2026年6月20日向本院移送审查起诉。本院受理后，依法告知被告人有权委托辩护人，依法讯问了被告人，听取了辩护人的意见，审查了全部案件材料。"),
    ("h1", "经依法审查查明："),
    ("body", "2025年3月至6月间，被告人陈某以其参与“星岚储能一期项目”、能够取得公司内部投资份额为由，向刘某、周某、郑某等人推介投资。陈某在沟通中称相关款项将用于项目设备采购和项目建设，并以项目退出、收益安排等内容促使三人付款。"),
    ("body", "2025年3月20日至2025年4月30日期间，刘某按照陈某要求向陈某指定账户支付首轮“星岚储能一期项目”投资款共计人民币240万元。"),
    ("body", "2025年4月12日至2025年4月22日期间，周某按照陈某要求为参与“星岚储能一期项目”支付投资款共计人民币180万元，其中部分款项由其配偶孙某代为转账。"),
    ("body", "2025年6月5日至2025年6月12日期间，郑某按照陈某要求为参与“星岚储能一期项目”支付投资款共计人民币160万元，其中部分款项转入陈某指定的马某账户。"),
    ("body", "2025年5月1日至2025年5月20日期间，刘某基于陈某关于“追加内部份额”的说明，再次支付款项共计人民币120万元，其中部分款项转入陈某指定的马某账户。"),
    ("body", "上述四项付款合计人民币700万元。经调取银行流水发现，相关款项进入陈某本人及马某账户后，部分流向设备供应商、证券账户、债权人、消费商户及亲友账户，并向刘某、周某、郑某账户转回部分资金。相关资金用途及转回款项性质，需结合在案其他证据综合认定。"),
    ("h1", "认定上述事实的证据如下："),
    ("body", "1. 被害人刘某、周某、郑某的陈述及其提供的微信聊天记录、转账记录；"),
    ("body", "2. 被告人陈某的供述与辩解；"),
    ("body", "3. 证人孙某、马某、高某、许某、韩某、赵某、马某甲等人的证言；"),
    ("body", "4. 刘某、周某、郑某、陈某、马某、孙某名下相关银行账户交易明细及银行回单；"),
    ("body", "5. 陈某名下证券资金账户的资金流水及交易记录；"),
    ("body", "6. 涉案项目相关电子文件、项目资料、设备采购材料及微信聊天记录电子数据；"),
    ("body", "7. 相关项目主体出具的情况说明、工商登记资料及设备供应商出具的证明材料；"),
    ("body", "8. 其他能够相互印证的书证、电子数据。"),
    ("body", "本院认为，被告人陈某以非法占有为目的，采用虚构事实、隐瞒真相的方法，骗取他人财物，数额特别巨大，其行为触犯《中华人民共和国刑法》第二百六十六条，犯罪事实清楚，证据确实、充分，应当以诈骗罪追究其刑事责任。根据《中华人民共和国刑事诉讼法》第一百七十六条之规定，提起公诉，请依法判处。"),
    ("thiszhi", "此致"),
    ("court", "北京市某区人民法院"),
    ("sign", "检察官：王某某"),
    ("sign", "2026年8月18日"),
    ("attach", "附：1. 被告人陈某现羁押于北京市某区看守所；"),
    ("attach", "2. 案卷材料及证据目录随案移送。"),
]

# ============ 03 刘某陈述 ============
liu = [
    ("title", "北京市公安局某分局"),
    ("title", "被害人陈述笔录"),
    ("subtitle", "被害人：刘某"),
    ("table", [
        ["询问时间", "2025年8月3日 09:30-11:20", "询问地点", "某分局询问室"],
        ["询问人", "民警李某", "记录人", "民警王某"],
        ["被询问人", "刘某，女，1983年生", "身份", "本案被害人"],
    ]),
]
liu += qa("你是如何认识陈某的？他对你如何介绍“星岚储能一期项目”？",
          "我于2025年3月经朋友介绍认识陈某。陈某说他正在参与“星岚储能一期项目”，项目已经进入设备采购阶段，有一部分内部投资份额，可以由熟人参与。我问过风险，他说项目真实、有设备和合作方，资金主要用于项目采购。")
liu += qa("你先后向陈某支付过多少投资款？",
          "我先后向陈某本人账户转了三笔钱，每笔80万元，共240万元。第一笔是2025年3月20日，第二笔是4月6日，第三笔是4月25日。陈某当时一直说这是一期项目投资款。")
liu += qa("2025年5月，陈某是否让你追加投资？",
          "到了5月初，陈某又说前面有投资人退出，空出120万元内部份额，让我追加。我于5月2日先给陈某本人账户转了60万元，5月12日又按陈某发来的账户给“马会计”转了60万元。陈某在聊天里把那个人叫“马会计”，我理解这个账户就是项目财务收款账户。")
liu += qa("你是否认识收到上述款项的账户实际持有人？",
          "后来我知道这个账户户名实际是马某。我本人没有见过马某，也不知道马某是不是项目公司的正式工作人员。")
liu += qa("你是否收到过转回的款项？",
          "7月以后，我收到过几笔钱。陈某本人给我转过30万元，摘要写“收益”；马某给我转过15万元，摘要写“退款”；陈某后来又给我转25万元，摘要写“周转”。陈某说这些都是本金和收益安排，但我当时理解其中至少有些钱是为了让我继续相信项目，我不能确认每一笔到底是什么性质。")
liu += [
    ("note", "说明：刘某称其并未因收到上述转回款项而与陈某重新结算投资份额，也未形成书面解除或清算协议。"),
    ("body", "被询问人核对意见：以上笔录我已阅读（听读），与我陈述一致。"),
    ("sign", "被询问人：刘某（签名处略）    2025年8月3日"),
]

# ============ 04 周某陈述 ============
zhou = [
    ("title", "北京市公安局某分局"),
    ("title", "被害人陈述笔录"),
    ("subtitle", "被害人：周某"),
    ("table", [
        ["询问时间", "2025年8月4日 09:30-10:50", "询问地点", "某分局询问室"],
        ["询问人", "民警李某", "记录人", "民警王某"],
        ["被询问人", "周某，女，1985年生", "身份", "本案被害人"],
    ]),
]
zhou += qa("你是何时、如何决定参与“星岚储能一期项目”的？",
           "2025年4月，陈某向我介绍“星岚储能一期项目”，说项目已经落实合作方，正在筹集设备采购资金。我决定投入180万元。")
zhou += qa("你如何支付投资款？其中有无他人代付？",
           "我本人分别于4月12日和4月22日各向陈某账户转了60万元。4月16日那笔60万元不是从我本人银行卡转的，是我让丈夫孙某从他的账户代我支付给陈某。我当时把陈某的收款账户发给孙某，并明确告诉他这笔是替我支付星岚项目投资款。")
zhou += qa("孙某与你之间是什么关系？该笔款项是否为其本人投资？",
           "我和孙某之间没有借贷或买卖关系，也不是孙某自己投资。他只是代我转账。")
zhou += qa("你是否收到过转回的款项？",
           "7月10日陈某向我转回25万元，摘要写“本金”。我没有和陈某就剩余款项做最终结算，也不能确认这25万元在法律上应当如何处理。")
zhou += [
    ("body", "被询问人核对意见：以上笔录我已阅读（听读），与我陈述一致。"),
    ("sign", "被询问人：周某（签名处略）    2025年8月4日"),
]

# ============ 05 郑某陈述 ============
zheng = [
    ("title", "北京市公安局某分局"),
    ("title", "被害人陈述笔录"),
    ("subtitle", "被害人：郑某"),
    ("table", [
        ["询问时间", "2025年8月5日 09:30-10:40", "询问地点", "某分局询问室"],
        ["询问人", "民警李某", "记录人", "民警王某"],
        ["被询问人", "郑某，女，1986年生", "身份", "本案被害人"],
    ]),
]
zheng += qa("你是如何了解“星岚储能一期项目”的？",
            "我于2025年6月听陈某介绍“星岚储能一期项目”。陈某说项目是真的，现阶段有内部额度，可以投入160万元。")
zheng += qa("你如何支付投资款？",
            "6月5日我向陈某本人账户转了80万元。6月12日，陈某又把一个户名为马某的账户发给我，让我把剩余80万元转过去。他说“这是马会计那边的项目归集账户”。")
zheng += qa("你是否认识马某？为何向其账户转账？",
            "我不认识马某，也不知道马某是不是项目会计。我之所以转账，是因为陈某明确让我转到该账户。")
zheng += qa("你是否收到过转回的款项？",
            "7月12日陈某向我转回20万元，摘要写“项目分红”。我没有收到项目公司的正式分红文件，也没有签署收益确认单。")
zheng += [
    ("body", "被询问人核对意见：以上笔录我已阅读（听读），与我陈述一致。"),
    ("sign", "被询问人：郑某（签名处略）    2025年8月5日"),
]

# ============ 06 证人证言汇编 ============
witness = [
    ("title", "北京市公安局某分局"),
    ("title", "证 人 证 言 汇 编"),
    ("subtitle", "GOLD_CASE_002｜孙某、马某、高某、许某、韩某、赵某、马某甲"),
    ("h1", "一、证人孙某证言"),
    ("table", [
        ["询问时间", "2025年8月6日 09:30-10:20", "询问地点", "某分局询问室"],
        ["证人", "孙某，男，1981年生", "与案件关系", "被害人周某之夫，代周某转账60万元"],
    ]),
]
witness += qa("你是否替周某转账？",
              "我是周某的丈夫。2025年4月16日，周某把一个收款账户发给我，让我代她转60万元，说是她参加陈某介绍的储能项目的投资款。我按她给的账户转了钱。")
witness += qa("这笔60万元是否为你本人投资或借款？",
              "这60万元不是我本人投资，也不是我借给陈某的钱。我只是替周某付款。")
witness += [
    ("h1", "二、证人马某证言"),
    ("table", [
        ["询问时间", "2025年8月8日 14:30-16:00", "询问地点", "某分局询问室"],
        ["证人", "马某，男，1984年生", "与案件关系", "账户A105持有人，代陈某收款、转款"],
    ]),
]
witness += qa("你是否为“星岚储能一期项目”的会计或项目公司员工？",
              "我认识陈某多年，但我不是“星岚储能一期项目”的会计，也不是项目公司的员工。陈某有时叫我帮忙收一下钱、再按他说的用途转出去。")
witness += qa("你的账户是否收到过涉案款项？",
              "2025年5月至6月，我的A105账户收到刘某60万元、郑某80万元，共140万元。陈某说是项目有关款项，让我先收着。之后我按他的要求支付过设备款、转给熟人，也给刘某退过15万元。")
witness += qa("你是否决定过这些资金的最终用途？",
              "我没有决定这些钱最终用于哪里，也没有以项目财务身份向刘某、郑某作过承诺。")
witness += [
    ("h1", "三、证人高某证言"),
    ("table", [
        ["询问时间", "2025年8月10日 10:00-10:40", "询问地点", "某分局询问室"],
        ["证人", "高某，男，1975年生", "与案件关系", "陈某债权人"],
    ]),
]
witness += qa("陈某向你转账的情况如何？",
              "陈某在2024年就欠我钱。2025年5月9日他转给我95万元，是偿还此前形成的借款，与后来所谓储能项目没有关系。")
witness += [
    ("h1", "四、证人许某证言"),
    ("table", [
        ["询问时间", "2025年8月10日 14:30-15:10", "询问地点", "某分局询问室"],
        ["证人", "许某，男，1978年生", "与案件关系", "陈某债权人"],
    ]),
]
witness += qa("陈某向你转账的情况如何？",
              "陈某此前欠我70万元。2025年5月18日和5月21日分别收到30万元、40万元，共70万元，都是归还旧债。债务形成时间早于2025年3月。")
witness += [
    ("h1", "五、证人韩某证言"),
    ("table", [
        ["询问时间", "2025年8月12日 09:30-11:00", "询问地点", "某分局询问室"],
        ["证人", "韩某，男，1980年生", "与案件关系", "星岚储能项目合作方人员"],
    ]),
]
witness += qa("陈某是否参与“星岚储能一期项目”？",
              "“星岚储能一期项目”确实存在，陈某曾参与前期资源对接，也介绍过部分设备供应商。")
witness += qa("项目公司是否授权陈某对外募集资金或出售“内部投资份额”？",
              "项目公司没有授权陈某向社会个人募集项目投资款，也没有所谓对外出售的“内部投资份额”。我不知道陈某向刘某、周某、郑某所说的固定退出或收益安排。")
witness += qa("项目设备采购能否证明刘某等人享有投资份额？",
              "项目确实发生过设备采购，但设备采购事实不能说明刘某等人拥有项目投资份额。")
witness += [
    ("h1", "六、证人赵某证言"),
    ("table", [
        ["询问时间", "2025年8月14日 09:30-10:30", "询问地点", "某分局询问室"],
        ["证人", "赵某，男，1972年生", "与案件关系", "设备供应商"],
    ]),
]
witness += qa("你方是否收到与储能设备有关的款项？",
              "我方在2025年4月至6月收到与储能设备有关的三笔款项，合计155万元，对应设备预付款及采购款。付款账户分别来自陈某账户及马某账户。")
witness += qa("你是否清楚上述款项的性质？",
              "我们只负责设备交易，不清楚资金是否属于刘某等人的投资款，也没有向刘某等人出具投资权益凭证。")
witness += [
    ("h1", "七、证人马某甲证言"),
    ("table", [
        ["询问时间", "2025年8月16日 10:00-10:30", "询问地点", "某分局询问室"],
        ["证人", "马某甲，男，1968年生", "与案件关系", "设备运输人员，外号“马师傅”"],
    ]),
]
witness += qa("你是否认识收取刘某、郑某款项的马某？",
              "我做设备运输，别人有时叫我“马师傅”。我不认识收取刘某、郑某款项的马某，也没有替陈某保管或归集投资款。")
witness += qa("你是否与设备供应商存在业务往来？",
              "我只在一次设备运输中与赵某一方有业务接触。")

# ============ 07 被告人供述与辩解 ============
defendant = [
    ("title", "北京市公安局某分局"),
    ("title", "被告人供述与辩解"),
    ("subtitle", "被告人：陈某｜三次讯问摘录"),
    ("h1", "第一次讯问"),
    ("table", [
        ["讯问时间", "2025年8月26日 16:20-18:30", "讯问地点", "某分局执法办案中心"],
        ["讯问人", "民警李某、民警张某", "记录人", "民警王某"],
    ]),
]
defendant += qa("你是否参与“星岚储能一期项目”？是否收取过刘某、周某、郑某的钱款？",
                "“星岚储能一期项目”是真实项目，我确实参与过项目对接。我承认刘某、周某、郑某给我或我指定账户转过钱。")
defendant += qa("三人分别向你支付了多少钱？",
                "刘某前后给了360万元，周某一方180万元，郑某160万元。我当时是想把这些钱用于项目周转，不是想骗他们。")
defendant += [
    ("h1", "第二次讯问"),
    ("table", [
        ["讯问时间", "2025年9月12日 09:10-11:40", "讯问地点", "某区看守所讯问室"],
        ["讯问人", "民警李某、民警张某", "记录人", "民警王某"],
    ]),
]
defendant += qa("马某是否为你所称的“马会计”？",
                "马某不是项目公司的正式会计，但我平时为了方便，会把他称作“马会计”。有部分款项我让刘某、郑某转到马某账户，马某按我的意思帮助收款或转款。")
defendant += qa("相关款项的用途是什么？",
                "我承认有一部分钱进入证券账户，也有部分用于还以前的债务和个人周转。我记得还旧债大约有210万元左右，但具体数额以银行流水为准。")
defendant += qa("“星岚储能一期项目”是否真实存在？",
                "确实有设备采购，项目不是完全虚构。")
defendant += [
    ("h1", "第三次讯问"),
    ("table", [
        ["讯问时间", "2025年10月15日 14:00-16:20", "讯问地点", "某区看守所讯问室"],
        ["讯问人", "民警李某、民警张某", "记录人", "民警王某"],
    ]),
]
defendant += qa("你是否向三名被害人转回过款项？",
                "我后来陆续给刘某、周某、郑某转过钱。我认为这些是退回本金或者收益安排的一部分。")
defendant += qa("你对本案指控如何辩解？",
                "我不同意说我一开始就想非法占有他们的钱。我承认具体资金用途没有完全按照当时向他们说的用途执行，但这不等于我承认自己实施诈骗。")
defendant += [("note", "说明：本材料保留被告人“承认客观事实但否认非法占有目的”的辩解结构，用于测试模型区分事实认定与主观要件评价的能力。")]

# ============ build all ============
jobs = [
    ("01_起诉书.docx", indictment),
    ("03_刘某陈述.docx", liu),
    ("04_周某陈述.docx", zhou),
    ("05_郑某陈述.docx", zheng),
    ("06_证人证言.docx", witness),
    ("07_被告人供述与辩解.docx", defendant),
]
for name, blocks in jobs:
    out = os.path.join(BUILD, name)
    build(blocks, out)
    print("built", name)
print("done")
