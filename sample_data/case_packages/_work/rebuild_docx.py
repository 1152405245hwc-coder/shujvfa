# -*- coding: utf-8 -*-
"""
Rebuild visible legal documents in GOLD_CASE_001 / GOLD_CASE_002 to
standard PRC legal-document typography, preserving text 100%.
Standard applied (公文/诉讼文书惯例):
  - A4, margins T3.7 B3.5 L2.8 R2.6 cm
  - Main title: SimHei(黑体) 22pt centered
  - Subtitle/文号: FangSong(仿宋) 16pt centered
  - Body: FangSong 16pt, justified, fixed 28pt line spacing, first-line indent 2 chars
  - H1: SimHei 16pt, indent 2 chars
  - Q/A prefix: SimHei 16pt; content FangSong 16pt
  - Signature lines: FangSong 16pt right aligned
  - Info tables: labels SimHei 14pt on #F2F2F2, values FangSong 14pt, 0.5pt borders
"""
import docx, os, re, json
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph

ROOT = "C:/Users/youxi/Documents/shujvfa/sample_data/case_packages"
BUILD = os.path.join(ROOT, "_work", "build")
os.makedirs(BUILD, exist_ok=True)

F_TITLE_EA = "黑体"
F_TITLE_ASCII = "SimHei"
F_BODY_EA = "仿宋"
F_BODY_ASCII = "Times New Roman"
F_NOTE_EA = "楷体"

HEADER_SPEC = {
    "GOLD_CASE_001/visible/documents/01_起诉书.docx": ["title", "title", "subtitle"],
    "GOLD_CASE_001/visible/documents/03_证人证言.docx": ["title", "title", "subtitle"],
    "GOLD_CASE_001/visible/documents/04_被告人供述与辩解.docx": ["title", "title", "subtitle"],
    "GOLD_CASE_001/visible/documents/05_被害人陈述.docx": ["title", "title", "subtitle"],
    "GOLD_CASE_002/visible/documents/01_indictment.docx": ["title", "subtitle"],
    "GOLD_CASE_002/visible/documents/03_victim_liu_statement.docx": ["title", "subtitle"],
    "GOLD_CASE_002/visible/documents/04_victim_zhou_statement.docx": ["title", "subtitle"],
    "GOLD_CASE_002/visible/documents/05_victim_zheng_statement.docx": ["title", "subtitle"],
    "GOLD_CASE_002/visible/documents/06_witness_statements.docx": ["title", "subtitle"],
    "GOLD_CASE_002/visible/documents/07_defendant_statements.docx": ["title", "subtitle"],
    "GOLD_CASE_002/visible/documents/08_chat_records.docx": ["title", "subtitle"],
    "GOLD_CASE_002/visible/documents/09_project_materials.docx": ["title", "subtitle"],
}

H1_PREFIXES = ("一、", "二、", "三、", "四、", "五、", "六、", "七、", "八、", "九、")
H1_EXACT = ("第一次讯问", "第二次讯问", "第三次讯问", "审查提示", "补充说明")


def set_run_font(run, ea, ascii_f, size, bold=False):
    run.font.name = ascii_f
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), ascii_f)
    rFonts.set(qn('w:hAnsi'), ascii_f)
    rFonts.set(qn('w:eastAsia'), ea)


def set_first_line_chars(p, chars=200, twips=640):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), str(chars))
    ind.set(qn('w:firstLine'), str(twips))


def clear_indent(p):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), '0')
    ind.set(qn('w:firstLine'), '0')
    ind.set(qn('w:leftChars'), '0')
    ind.set(qn('w:left'), '0')


def para_props(p, align=None, line_exact=None, line_mult=None, before=0, after=0):
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if line_exact:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line_exact)
    elif line_mult:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line_mult
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def add_para(doc, text, kind, after=0):
    p = doc.add_paragraph()
    if kind == "title":
        para_props(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_mult=1.0, before=0, after=after)
        r = p.add_run(text)
        set_run_font(r, F_TITLE_EA, F_TITLE_ASCII, 22)
    elif kind == "subtitle":
        para_props(p, align=WD_ALIGN_PARAGRAPH.CENTER, line_mult=1.0, before=0, after=after)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "h1":
        para_props(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_exact=28, before=0, after=0)
        set_first_line_chars(p)
        r = p.add_run(text)
        set_run_font(r, F_TITLE_EA, F_TITLE_ASCII, 16)
    elif kind == "qa":
        para_props(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_exact=28, before=0, after=0)
        set_first_line_chars(p)
        m = re.match(r'^(问|答)：', text)
        if m:
            r1 = p.add_run(text[:m.end()])
            set_run_font(r1, F_TITLE_EA, F_TITLE_ASCII, 16)
            r2 = p.add_run(text[m.end():])
            set_run_font(r2, F_BODY_EA, F_BODY_ASCII, 16)
        else:
            r = p.add_run(text)
            set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "body":
        para_props(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_exact=28, before=0, after=0)
        set_first_line_chars(p)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "thiszhi":
        para_props(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_exact=28, before=0, after=0)
        set_first_line_chars(p)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "court":
        para_props(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_exact=28, before=0, after=0)
        clear_indent(p)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "sign":
        para_props(p, align=WD_ALIGN_PARAGRAPH.RIGHT, line_exact=28, before=0, after=0)
        clear_indent(p)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "attach":
        para_props(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_exact=28, before=0, after=0)
        clear_indent(p)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    elif kind == "note":
        para_props(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_exact=24, before=6, after=0)
        set_first_line_chars(p, 200, 560)
        r = p.add_run(text)
        set_run_font(r, F_NOTE_EA, F_BODY_ASCII, 14)
    elif kind == "ts":
        para_props(p, align=WD_ALIGN_PARAGRAPH.LEFT, line_exact=24, before=6, after=0)
        clear_indent(p)
        r = p.add_run(text)
        set_run_font(r, F_TITLE_EA, F_TITLE_ASCII, 14)
    else:
        para_props(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, line_exact=28)
        set_first_line_chars(p)
        r = p.add_run(text)
        set_run_font(r, F_BODY_EA, F_BODY_ASCII, 16)
    return p


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), '000000')
        borders.append(el)
    tblPr.append(borders)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def vcenter_cell(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    v = OxmlElement('w:vAlign')
    v.set(qn('w:val'), 'center')
    tcPr.append(v)


def add_table(doc, rows_data, widths):
    n_rows = len(rows_data)
    n_cols = len(rows_data[0])
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for ri, row in enumerate(rows_data):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.width = Cm(widths[ci])
            vcenter_cell(cell)
            p = cell.paragraphs[0]
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.space_before = Pt(1)
            pf.space_after = Pt(1)
            clear_indent(p)
            is_label = (ci % 2 == 0)
            is_short = len(val) <= 8
            if is_label:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shade_cell(cell, "F2F2F2")
                r = p.add_run(val)
                set_run_font(r, F_TITLE_EA, F_TITLE_ASCII, 14)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_short else WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(val)
                set_run_font(r, F_BODY_EA, F_BODY_ASCII, 14)
    return table


def classify(text):
    if not text:
        return "skip"
    if text.startswith(("问：", "答：")):
        return "qa"
    if text.startswith(H1_PREFIXES):
        return "h1"
    if text in H1_EXACT:
        return "h1"
    if text.startswith("证人：") or text.startswith("经依法审查查明"):
        return "h1"
    if text.startswith("此致"):
        return "thiszhi"
    if text.startswith("北京市某区人民法院"):
        return "court"
    if text.startswith("检察官：") or text.startswith("被询问人：") or re.match(r'^\d{4}年\d+月\d+日$', text):
        return "sign"
    if text.startswith("附："):
        return "attach"
    if text.startswith("说明："):
        return "note"
    if re.match(r'^20\d{2}-\d{2}-\d{2} \d{2}:\d{2}', text):
        return "ts"
    return "body"


def build_docx(src_path, spec, out_path):
    src = Document(src_path)
    doc = Document()
    # Normal style base
    st = doc.styles['Normal']
    st.font.name = F_BODY_ASCII
    st.font.size = Pt(16)
    rPr = st.element.get_or_add_rPr()
    rF = rPr.find(qn('w:rFonts'))
    if rF is None:
        rF = OxmlElement('w:rFonts')
        rPr.append(rF)
    rF.set(qn('w:ascii'), F_BODY_ASCII)
    rF.set(qn('w:hAnsi'), F_BODY_ASCII)
    rF.set(qn('w:eastAsia'), F_BODY_EA)
    # page setup
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(3.7)
    sec.bottom_margin = Cm(3.5)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)
    # content
    body = src.element.body
    para_idx = 0
    blocks = []  # (kind, payload)
    for child in body.iterchildren():
        if child.tag.endswith('}p'):
            p = Paragraph(child, src)
            txt = p.text
            if para_idx < len(spec):
                kind = spec[para_idx]
                para_idx += 1
            else:
                kind = classify(txt)
            blocks.append((kind, txt))
        elif child.tag.endswith('}tbl'):
            t = DocxTable(child, src)
            rows_data = [[c.text for c in row.cells] for row in t.rows]
            blocks.append(("table", rows_data))
    # emit
    for idx, (kind, payload) in enumerate(blocks):
        if kind == "skip":
            continue
        if kind == "table":
            n_cols = len(payload[0])
            if n_cols == 2:
                widths = [2.6, 13.0]
            elif n_cols == 4:
                widths = [2.6, 5.2, 2.6, 5.2]
            else:
                w = 15.6 / n_cols
                widths = [w] * n_cols
            add_table(doc, payload, widths)
            # spacing paragraph after table
            sp = doc.add_paragraph()
            para_props(sp, line_exact=14)
        else:
            after = 0
            if kind == "title":
                # check if it's the last title line
                nxt = blocks[idx + 1][0] if idx + 1 < len(blocks) else None
                after = 6 if nxt == "title" else 18
            elif kind == "subtitle":
                after = 12
            add_para(doc, payload, kind, after=after)
    doc.save(out_path)
    return blocks


def extract_texts(path):
    d = Document(path)
    out_p = []
    out_t = []
    for child in d.element.body.iterchildren():
        if child.tag.endswith('}p'):
            p = Paragraph(child, d)
            if p.text.strip():
                out_p.append(p.text)
        elif child.tag.endswith('}tbl'):
            t = DocxTable(child, d)
            for row in t.rows:
                for c in row.cells:
                    out_t.append(c.text)
    return out_p, out_t


report = []
for rel, spec in HEADER_SPEC.items():
    src = os.path.join(ROOT, rel.replace("/", os.sep))
    name = os.path.basename(rel)
    out = os.path.join(BUILD, name)
    blocks = build_docx(src, spec, out)
    src_p, src_t = extract_texts(src)
    out_p, out_t = extract_texts(out)
    ok_p = src_p == out_p
    ok_t = src_t == out_t
    status = "OK" if (ok_p and ok_t) else "MISMATCH"
    report.append((rel, status, len(src_p), len(out_p), len(src_t), len(out_t)))
    if not ok_p:
        for a, b in zip(src_p, out_p):
            if a != b:
                report.append(("  para diff", a[:50], b[:50]))
    if not ok_t:
        for a, b in zip(src_t, out_t):
            if a != b:
                report.append(("  tbl diff", a[:50], b[:50]))

with open(os.path.join(ROOT, "_work", "rebuild_report.txt"), "w", encoding="utf-8") as fh:
    for line in report:
        fh.write(str(line) + "\n")
for line in report:
    print(line)
