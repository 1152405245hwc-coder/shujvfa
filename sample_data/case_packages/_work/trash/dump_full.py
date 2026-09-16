# -*- coding: utf-8 -*-
"""Dump full text content of all docx (for faithful rebuild) + inspect xlsx."""
import docx, os, glob, json
from docx.table import Table
from docx.text.paragraph import Paragraph

def dump_full_text(path):
    d = docx.Document(path)
    lines = []
    body = d.element.body
    for child in body.iterchildren():
        if child.tag.endswith('}p'):
            p = Paragraph(child, d)
            lines.append({"t": "p", "text": p.text})
        elif child.tag.endswith('}tbl'):
            t = Table(child, d)
            rows = []
            for row in t.rows:
                rows.append([c.text for c in row.cells])
            lines.append({"t": "tbl", "rows": rows})
    return lines

root = "C:/Users/youxi/Documents/shujvfa/sample_data/case_packages"
out = {}
for f in sorted(glob.glob(os.path.join(root, "GOLD_CASE_00?", "**", "*.docx"), recursive=True)):
    key = f.replace("\\", "/")
    out[key] = dump_full_text(f)

with open(os.path.join(root, "_work", "docx_fulltext.json"), "w", encoding="utf-8") as fh:
    json.dump(out, fh, ensure_ascii=False, indent=1)
print("docx files:", len(out))

# inspect xlsx
try:
    import openpyxl
    for f in sorted(glob.glob(os.path.join(root, "GOLD_CASE_00?", "**", "*.xlsx"), recursive=True)):
        wb = openpyxl.load_workbook(f, data_only=True)
        info = {"sheets": []}
        for ws in wb.worksheets:
            s = {"name": ws.title, "dims": ws.dimensions, "max_row": ws.max_row, "max_col": ws.max_column,
                 "merged": [str(m) for m in ws.merged_cells.ranges][:10],
                 "freeze": str(ws.freeze_panes),
                 "sample_rows": []}
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i < 8:
                    s["sample_rows"].append([str(c)[:20] if c is not None else "" for c in row])
            info["sheets"].append(s)
        out[f.replace("\\","/")] = info
    with open(os.path.join(root, "_work", "xlsx_inspect.json"), "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
    print("xlsx inspected")
except Exception as e:
    print("openpyxl err:", e)
