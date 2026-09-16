# -*- coding: utf-8 -*-
"""Final verification of reformatted packages against backups."""
import docx, os, json, glob
import openpyxl
from docx.table import Table as DTable
from docx.text.paragraph import Paragraph

ROOT = "C:/Users/youxi/Documents/shujvfa/sample_data/case_packages"
BK = os.path.join(ROOT, "_work", "backup")

def texts(path):
    d = docx.Document(path)
    ps, ts = [], []
    for child in d.element.body.iterchildren():
        if child.tag.endswith('}p'):
            p = Paragraph(child, d)
            if p.text.strip():
                ps.append(p.text)
        elif child.tag.endswith('}tbl'):
            t = DTable(child, d)
            for row in t.rows:
                for c in row.cells:
                    ts.append(c.text)
    return ps, ts

def fmt_check(path):
    d = docx.Document(path)
    sec = d.sections[0]
    info = {
        "page_cm": (round(sec.page_width.cm, 2), round(sec.page_height.cm, 2)),
        "margins_cm": (round(sec.top_margin.cm, 2), round(sec.bottom_margin.cm, 2),
                       round(sec.left_margin.cm, 2), round(sec.right_margin.cm, 2)),
        "paras": []
    }
    n = 0
    for child in d.element.body.iterchildren():
        if not child.tag.endswith('}p'):
            continue
        p = Paragraph(child, d)
        if not p.text.strip():
            continue
        n += 1
        if n > 60:
            break
        r = p.runs[0] if p.runs else None
        if r is None:
            continue
        rPr = r._element.rPr
        ea = None
        if rPr is not None and rPr.rFonts is not None:
            ea = rPr.rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia')
        sz = r.font.size.pt if r.font.size else None
        ls = p.paragraph_format.line_spacing
        ls_pt = ls.pt if hasattr(ls, "pt") and ls else None
        ind = p._p.pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ind') if p._p.pPr is not None else None
        flc = ind.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}firstLineChars') if ind is not None else None
        if n <= 3 or p.text.startswith(("问：", "答：", "一、", "经依法审查")):
            info["paras"].append((p.text[:22], ea, sz, ls_pt, flc))
    return info

def xlsx_snapshot(path):
    wb = openpyxl.load_workbook(path)
    return {ws.title: ([c.value for row in ws.iter_rows() for c in row],
                       sorted(str(m) for m in ws.merged_cells.ranges)) for ws in wb.worksheets}

report = []
# docx visible files
mapping = {
    "GOLD_CASE_001/visible/documents/01_起诉书.docx": "GOLD_CASE_001/visible/documents/01_起诉书.docx",
    "GOLD_CASE_001/visible/documents/03_证人证言.docx": "GOLD_CASE_001/visible/documents/03_证人证言.docx",
    "GOLD_CASE_001/visible/documents/04_被告人供述与辩解.docx": "GOLD_CASE_001/visible/documents/04_被告人供述与辩解.docx",
    "GOLD_CASE_001/visible/documents/05_被害人陈述.docx": "GOLD_CASE_001/visible/documents/05_被害人陈述.docx",
    "GOLD_CASE_002/visible/documents/01_起诉书.docx": "GOLD_CASE_002/visible/documents/01_indictment.docx",
    "GOLD_CASE_002/visible/documents/03_刘某陈述.docx": "GOLD_CASE_002/visible/documents/03_victim_liu_statement.docx",
    "GOLD_CASE_002/visible/documents/04_周某陈述.docx": "GOLD_CASE_002/visible/documents/04_victim_zhou_statement.docx",
    "GOLD_CASE_002/visible/documents/05_郑某陈述.docx": "GOLD_CASE_002/visible/documents/05_victim_zheng_statement.docx",
    "GOLD_CASE_002/visible/documents/06_证人证言.docx": "GOLD_CASE_002/visible/documents/06_witness_statements.docx",
    "GOLD_CASE_002/visible/documents/07_被告人供述与辩解.docx": "GOLD_CASE_002/visible/documents/07_defendant_statements.docx",
    "GOLD_CASE_002/visible/documents/08_聊天记录摘录.docx": "GOLD_CASE_002/visible/documents/08_chat_records.docx",
    "GOLD_CASE_002/visible/documents/09_项目资料.docx": "GOLD_CASE_002/visible/documents/09_project_materials.docx",
}
for new_rel, old_rel in mapping.items():
    new_p = os.path.join(ROOT, new_rel.replace("/", os.sep))
    old_p = os.path.join(BK, old_rel.replace("/", os.sep))
    np_, nt_ = texts(new_p)
    op_, ot_ = texts(old_p)
    same = (np_ == op_) and (nt_ == ot_)
    fmt = fmt_check(new_p)
    report.append(("DOCX", new_rel, "TEXT_OK" if same else "TEXT_DIFF", fmt["page_cm"], fmt["margins_cm"], len(fmt["paras"])))
    for row in fmt["paras"]:
        report.append(("  fmt",) + row)

xlsx_map = {
    "GOLD_CASE_001/visible/bank/02_银行流水账单.xlsx": "GOLD_CASE_001/visible/bank/02_银行流水账单.xlsx",
    "GOLD_CASE_002/visible/bank/02_银行流水账单.xlsx": "GOLD_CASE_002/visible/bank/02_bank_statements.xlsx",
}
for new_rel, old_rel in xlsx_map.items():
    new_p = os.path.join(ROOT, new_rel.replace("/", os.sep))
    old_p = os.path.join(BK, old_rel.replace("/", os.sep))
    nw = xlsx_snapshot(new_p)
    ow = xlsx_snapshot(old_p)
    ok = set(nw) == set(ow)
    val_ok = all(nw[k][0] == ow[k][0] for k in ow)
    mrg_ok = all(nw[k][1] == ow[k][1] for k in ow)
    report.append(("XLSX", new_rel, "OK" if (ok and val_ok and mrg_ok) else "DIFF",
                   "sheets_same=" + str(ok), "values_same=" + str(val_ok), "merges_same=" + str(mrg_ok)))

with open(os.path.join(ROOT, "_work", "final_verify.txt"), "w", encoding="utf-8") as fh:
    for line in report:
        fh.write(" | ".join(str(x) for x in line) + "\n")
for line in report:
    print(" | ".join(str(x) for x in line))
