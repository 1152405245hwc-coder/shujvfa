# -*- coding: utf-8 -*-
"""
Restyle visible bank-statement workbooks as standard evidence materials.
Preserves all cell values, sheet names and merged ranges.
"""
import openpyxl, os, copy
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

ROOT = "C:/Users/youxi/Documents/shujvfa/sample_data/case_packages"
BUILD = os.path.join(ROOT, "_work", "build")

thin = Side(style="thin", color="000000")
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)
fill_gray = PatternFill(fill_type="solid", fgColor="F2F2F2")
F_TITLE = Font(name="黑体", size=16, bold=True)
F_INFO = Font(name="仿宋", size=12)
F_INFO_BOLD = Font(name="黑体", size=12, bold=True)
F_HEAD = Font(name="黑体", size=12, bold=True)
F_DATA = Font(name="仿宋", size=11)
AL_C = Alignment(horizontal="center", vertical="center", wrap_text=False)
AL_L = Alignment(horizontal="left", vertical="center", wrap_text=False)
AL_R = Alignment(horizontal="right", vertical="center")


def find_header_row(ws):
    for r in range(1, min(ws.max_row, 8) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and str(v).strip() == "交易时间":
                return r
    return None


def style_evidence_sheet(ws):
    hdr = find_header_row(ws)
    if hdr is None:
        # info-only sheet (e.g., 00_说明)
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                cell = ws.cell(row=r, column=c)
                if cell.value is None:
                    continue
                if r == 1:
                    cell.font = F_TITLE
                else:
                    cell.font = F_INFO
                cell.alignment = AL_L
        ws.freeze_panes = None
        # column widths
        for c in range(1, ws.max_column + 1):
            letter = get_column_letter(c)
            ws.column_dimensions[letter].width = 22
        return
    # determine money column
    money_col = None
    for c in range(1, ws.max_column + 1):
        if str(ws.cell(row=hdr, column=c).value).strip() == "金额":
            money_col = c
            break
    # title/info rows above header
    for r in range(1, hdr):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            if cell.value is None:
                continue
            txt = str(cell.value)
            if txt in ("户名", "账号", "账户ID") or (hdr - r == 1):
                cell.font = F_INFO_BOLD
            else:
                cell.font = F_INFO
            cell.alignment = AL_L
    ws.row_dimensions[hdr].height = 24
    # header row
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=hdr, column=c)
        cell.font = F_HEAD
        cell.fill = fill_gray
        cell.alignment = AL_C
        cell.border = border_all
    # data rows
    for r in range(hdr + 1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = F_DATA
            cell.border = border_all
            if c == money_col:
                cell.alignment = AL_R
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = AL_L
    # freeze below header
    ws.freeze_panes = ws.cell(row=hdr + 1, column=1)
    # column widths by content (capped)
    for c in range(1, ws.max_column + 1):
        letter = get_column_letter(c)
        maxlen = 10
        for r in range(hdr, min(ws.max_row, hdr + 20) + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                ln = len(str(v))
                if ln > maxlen:
                    maxlen = ln
        ws.column_dimensions[letter].width = min(maxlen + 4, 36)
    # page setup A4 portrait fit width
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def snapshot(ws):
    rows = []
    for row in ws.iter_rows():
        rows.append([c.value for c in row])
    return rows


workbooks = [
    ("GOLD_CASE_001/visible/bank/02_银行流水账单.xlsx", "02_银行流水账单_G001.xlsx"),
    ("GOLD_CASE_002/visible/bank/02_bank_statements.xlsx", "02_银行流水账单_G002.xlsx"),
]

report = []
for rel, outname in workbooks:
    src = os.path.join(ROOT, rel.replace("/", os.sep))
    wb = openpyxl.load_workbook(src)
    before = {ws.title: snapshot(ws) for ws in wb.worksheets}
    merges_before = {ws.title: sorted(str(m) for m in ws.merged_cells.ranges) for ws in wb.worksheets}
    for ws in wb.worksheets:
        style_evidence_sheet(ws)
    out = os.path.join(BUILD, outname)
    wb.save(out)
    # verify
    wb2 = openpyxl.load_workbook(out)
    ok = True
    for ws in wb2.worksheets:
        if ws.title not in before:
            ok = False
            continue
        if snapshot(ws) != before[ws.title]:
            ok = False
            report.append(("value diff", ws.title))
        m2 = sorted(str(m) for m in ws.merged_cells.ranges)
        if m2 != merges_before[ws.title]:
            ok = False
            report.append(("merge diff", ws.title))
    report.append((outname, "OK" if ok else "DIFF"))

with open(os.path.join(ROOT, "_work", "xlsx_report.txt"), "w", encoding="utf-8") as fh:
    for line in report:
        fh.write(str(line) + "\n")
for line in report:
    print(line)
